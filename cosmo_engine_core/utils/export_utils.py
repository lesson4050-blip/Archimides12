import json
import os
import aiohttp
from typing import Literal, List
import uuid
from fastapi import HTTPException
from pathvalidate import sanitize_filename

from models.presentation_and_path import PresentationAndPath
from models.sql.presentation import PresentationModel
from models.sql.slide import SlideModel
from services.direct_pptx_builder import DirectPptxBuilder
from services.pptx_themes import guess_theme
from services.temp_file_service import TEMP_FILE_SERVICE
from utils.asset_directory_utils import get_exports_directory


import asyncio
from pptx import Presentation
from pptx.util import Inches, Emu

async def export_presentation(
    presentation: PresentationModel, slides: List[SlideModel], export_as: Literal["pptx", "pdf"]
) -> PresentationAndPath:
    if export_as == "pptx":
        title = presentation.title or ""
        export_directory = get_exports_directory()
        safe_title = sanitize_filename(title or str(uuid.uuid4()))
        pptx_path = os.path.join(export_directory, f"{safe_title}.pptx")
        
        # 1. Use Puppeteer to capture Next.js slides as images
        temp_dir = os.path.abspath(TEMP_FILE_SERVICE.create_temp_dir())
        capture_script = r"d:\Cosmo\cosmo_artist\capture_slides.cjs"
        
        print(f"Running puppeteer capture for {presentation.id}")
        proc = await asyncio.create_subprocess_exec(
            "node", capture_script, str(presentation.id), temp_dir,
            stdout=asyncio.subprocess.PIPE,
            stderr=asyncio.subprocess.PIPE,
            cwd=r"d:\Cosmo\cosmo_artist"
        )
        stdout, stderr = await proc.communicate()
        print(stdout.decode())
        if stderr:
            print(f"Puppeteer stderr: {stderr.decode()}")
            
        # 2. Build PPTX using the captured images
        prs = Presentation()
        # Widescreen 16:9
        prs.slide_width = Emu(12192000)
        prs.slide_height = Emu(6858000)
        blank_slide_layout = prs.slide_layouts[6]
        
        # Get sorted screenshots
        images = []
        if os.path.exists(temp_dir):
            for file in os.listdir(temp_dir):
                if file.endswith('.png'):
                    images.append(file)
                    
        print(f"DEBUG: Found {len(images)} images in {temp_dir}")
                    
        # Sort by slide index (e.g., slide_0.png)
        if images:
            images.sort(key=lambda x: int(x.split('_')[1].split('.')[0]))
        
        for idx, img_file in enumerate(images):
            slide = prs.slides.add_slide(blank_slide_layout)
            img_path = os.path.join(temp_dir, img_file)
            print(f"DEBUG: Adding image to slide {idx}: {img_path}")
            
            # Full bleed background image
            slide.shapes.add_picture(img_path, 0, 0, prs.slide_width, prs.slide_height)
            
            # Add speaker notes if available
            if idx < len(slides):
                speaker_note = getattr(slides[idx], 'speaker_note', '') or getattr(slides[idx], 'speaker_notes', '') or ''
                if speaker_note:
                    slide.notes_slide.notes_text_frame.text = speaker_note
                    
        prs.save(pptx_path)

        return PresentationAndPath(
            presentation_id=presentation.id,
            path=pptx_path,
        )
    else:
        # PDF export: still uses Artist for now
        async with aiohttp.ClientSession() as session:
            artist_url = os.environ.get("ARTIST_URL", "http://localhost:3005")
            async with session.post(
                f"{artist_url}/api/export-as-pdf",
                json={
                    "id": str(presentation.id),
                    "title": sanitize_filename(presentation.title or str(uuid.uuid4())),
                },
            ) as response:
                response_json = await response.json()

        return PresentationAndPath(
            presentation_id=presentation.id,
            path=response_json.get("path") or "",
        )
