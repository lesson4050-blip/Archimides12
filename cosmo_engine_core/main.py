from fastapi import FastAPI
from pydantic import BaseModel
from fastapi.responses import Response

app = FastAPI()

@app.get("/api/health")
def health():
    return {"status": "ok"}

class OutlineReq(BaseModel):
    prompt: str
    slide_count: int
    language: str

@app.post("/api/v1/ppt/generate-outline")
def gen_outline(req: OutlineReq):
    return {"id": "mock_presentation_id"}

class FromOutlineReq(BaseModel):
    title: str
    outline: list
    theme: str
    language: str

@app.post("/api/v1/ppt/create-from-outline")
def create_outline(req: FromOutlineReq):
    return {"id": "mock_presentation_id"}

class GenPresReq(BaseModel):
    id: str
    theme: str
    fetch_images: bool

@app.post("/api/v1/ppt/generate-presentation")
def gen_pres(req: GenPresReq):
    return {"status": "ok"}

@app.get("/api/v1/ppt/download/{presentation_id}")
def download(presentation_id: str):
    try:
        from pptx import Presentation
    except ImportError:
        return Response(content=b"mock_pptx_content", media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation")
    
    from io import BytesIO
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "COSMO Presentation Mock"
    out = BytesIO()
    prs.save(out)
    return Response(
        content=out.getvalue(), 
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )
